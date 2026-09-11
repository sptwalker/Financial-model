# -*- coding: utf-8 -*-
"""工资夹具与解析器测试（隐私数据迁出后的回归护栏）

背景：真实工资表含个人信息，已移出仓库。这组测试确保：
1. 脱敏夹具存在、可生成，且金额与财务确认口径一致；
2. 解析逻辑只依赖夹具，任何指向 docs/ 下真实工资表的回归都会被挡下；
3. 夹具确实不含隐私列（姓名/身份证），防止有人把真实文件复制回来。
"""
import sys
from decimal import Decimal
from pathlib import Path

import pytest

BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))

from app.services.payroll import (  # noqa: E402
    FIXTURE,
    HOUSING_FUND_TOTAL_YUAN,
    SOCIAL_INSURANCE_TOTAL_YUAN,
    PayrollFixtureMissing,
    fixture_path,
    payroll_headcount,
    payroll_total_payable,
)

# 财务确认口径：2026-07 应付工资合计（元）
EXPECTED_PAYABLE_YUAN = Decimal("1741605.74")


def test_fixture_exists_and_is_readable():
    assert FIXTURE.exists(), (
        f"工资夹具缺失：{FIXTURE}\n"
        "生成方式：python -m scripts.build_payroll_fixture"
    )
    assert fixture_path() == FIXTURE


def test_payroll_total_matches_confirmed_caliber():
    """应付工资合计 = 174.160574 万元（对账基线的唯一来源）"""
    assert payroll_total_payable() == EXPECTED_PAYABLE_YUAN
    assert payroll_total_payable() / Decimal("10000") == Decimal("174.160574")


def test_headcount_is_aggregate_only():
    """在册人数只暴露总数，不含个人维度"""
    assert payroll_headcount() == 78


def test_social_insurance_and_housing_fund_totals():
    """社保/公积金合计为脱敏结论常量（明细表已移出仓库）"""
    assert SOCIAL_INSURANCE_TOTAL_YUAN == Decimal("523749.64")
    assert HOUSING_FUND_TOTAL_YUAN == Decimal("131083.10")


def test_fixture_contains_no_personal_columns():
    """夹具不得出现姓名/身份证类列 —— 防止真实文件被误复制回仓库"""
    import openpyxl

    wb = openpyxl.load_workbook(FIXTURE, data_only=True)
    banned = ("姓名", "身份证", "员工", "银行卡", "手机")
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is None:
                    continue
                text = str(cell)
                for word in banned:
                    assert word not in text, f"夹具出现隐私字段「{word}」：{text!r}"
        # 夹具应只有表头 + 总计行两行数据
        assert ws.max_row <= 2, f"夹具 {ws.title} 行数 {ws.max_row} 超出汇总范围"


def test_repo_no_longer_tracks_payroll_source():
    """仓库运行时不依赖 docs/ 下的真实工资表"""
    import subprocess

    tracked = subprocess.run(
        ["git", "ls-files", "docs/"],
        cwd=BACKEND_DIR.parent, capture_output=True, text=True, check=True,
    ).stdout
    assert "工资表" not in tracked, "真实工资表仍在 git 索引中"
    assert "社保" not in tracked, "社保明细仍在 git 索引中"
    assert "公积金" not in tracked, "公积金清单仍在 git 索引中"


def test_missing_fixture_raises_clear_error(tmp_path):
    """夹具缺失时必须显式报错，不能静默返回 0（否则现金链被污染）"""
    with pytest.raises(PayrollFixtureMissing) as exc:
        payroll_total_payable(tmp_path / "不存在.xlsx")
    assert "build_payroll_fixture" in str(exc.value)


def test_seed_and_reconcile_share_one_parser():
    """种子与对账必须走同一份解析逻辑，避免口径再次分叉"""
    from app.engine import reconcile
    from scripts import seed_from_excel

    assert reconcile.payroll_total_payable() == payroll_total_payable() / Decimal("10000")
    assert seed_from_excel.payroll_total_payable() == payroll_total_payable() / Decimal("10000")
    assert reconcile.PAYROLL_XLSX == FIXTURE
    assert seed_from_excel.PAYROLL_XLSX == FIXTURE
