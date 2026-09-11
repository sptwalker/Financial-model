#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成脱敏工资表夹具（backend/fixtures/payroll_2026-07.xlsx）

背景：真实的 `2026年7月创想悦动工资表.xlsx` 含员工姓名、身份证号、逐人工资明细等
个人信息，已移出仓库（见 .gitignore 与 docs/验收报告.md §6）。种子与对账只需要其中
**一个汇总结论**——「汇总」sheet 总计行的应付工资合计——因此这里生成一份结构相同、
但只保留总计行的脱敏替身。

夹具保留真实合计金额，因为该数字本身是财务口径结论、不含个人信息；数量级与真实值
一致才能保证对账基线的相对误差校验仍然有效。

真实文件准备就绪后，重建夹具的命令：

    python -m scripts.build_payroll_fixture --from-docs

（该开关只在本地、真实文件存在时可用；CI 与生产走默认的脱敏数据。）

字段口径见 backend/app/services/payroll.py。
"""
import argparse
import sys
from decimal import Decimal
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND_DIR))

OUT = BACKEND_DIR / "fixtures" / "payroll_2026-07.xlsx"

# 「汇总」sheet 表头（与真实工资表一致；解析器按列名定位，不依赖列序）
HEADER = [
    "分类", "部门", "人数", "基本工资", "绩效工资", "岗位津贴", "节日福利",
    "通讯补贴", "考勤扣款", "其他扣款", "加班费", "绩效奖金", "其他奖金",
    "伙食补贴", "住房补贴", "司机补贴", "专利奖", "其他补发", "补发合计",
    "应付工资", "借款逾期扣款", "房租水电", "房租水电1", "其他代扣款", "扣款合计",
    "养老保险", "医疗保险", "失业保险", "住房公积金", "所得税", "实发工资",
]

# 2026-07 汇总总计行（真实合计，已脱敏：无部门/个人维度）
TOTAL_ROW = [
    "总计", None, 78, 1790300, 0, 0, 0, 0, 49494.26, 0, 0, 0, 0, 0, 0,
    0, 0, 800, 800, 1741605.74, 0, 0, 0, -699.99, -699.99,
    119278.4, 31611.92, 3371.42, 131083.1, 126995.41, 1329965.48,
]


def build(out: Path = OUT) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws.append(HEADER)
    ws.append(TOTAL_ROW)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def rebuild_from_real() -> Path:
    """从 docs/ 下的真实工资表重建夹具（仅本地、真实文件存在时可用）。"""
    import openpyxl

    real = BACKEND_DIR.parent / "docs" / "2026年7月创想悦动工资表.xlsx"
    if not real.exists():
        raise SystemExit(f"真实工资表不存在：{real}\n"
                         "本仓库已将个人隐私文件移出，此开关仅供本地重建夹具使用。")
    src = openpyxl.load_workbook(real, data_only=True)["汇总"]
    hdr, total = None, None
    for row in src.iter_rows(values_only=True):
        if hdr is None and row and "应付工资" in [str(x) for x in row if x is not None]:
            hdr = list(row)
            continue
        if hdr is not None and row and row[0] == "总计":
            total = list(row)
            break
    if total is None:
        raise SystemExit("真实工资表未找到「总计」行")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws.append(hdr)
    ws.append(total)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    # 自检：夹具应付工资与真实值一致
    payable_col = hdr.index("应付工资")
    assert Decimal(str(total[payable_col])) == _fixture_payable(), "重建结果与预期不一致"
    return OUT


def _fixture_payable() -> Decimal:
    return Decimal(str(TOTAL_ROW[HEADER.index("应付工资")]))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="生成脱敏工资表夹具")
    ap.add_argument("--from-docs", action="store_true",
                    help="从 docs/ 下真实工资表重建（仅本地可用，含个人信息，勿提交）")
    args = ap.parse_args(argv)
    path = rebuild_from_real() if args.from_docs else build()
    print(f"已生成 {path}")
    print(f"应付工资合计 = {_fixture_payable() / 10000} 万元")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
