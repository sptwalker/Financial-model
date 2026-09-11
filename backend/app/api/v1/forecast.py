# -*- coding: utf-8 -*-
"""预测 API（阶段 3）：历史累积 → 运行预测 → 接受回填 → 回测/运行记录

历史唯一来源 sales_actuals；run 无副作用；apply 回填 qty 并 recalc（×1 或 ×3）。
错误映射照抄 scenarios.py（no_version/scenario_not_found→404，冲突→409）。
"""
import json
from datetime import datetime
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session

from app.api.deps import get_current_user
from app.core.permissions import PermissionChecker
from app.db.session import get_db
from app.models.forecast import SalesActual, ForecastRun
from app.schemas.forecast import ActualBatch, RunRequest, ApplyRequest
from app.services import forecast_service as fs
from app.services.operation_log_service import OperationLogService

router = APIRouter(prefix="/forecast", tags=["forecast"])

_CHANNELS = {"online", "offline"}


# ---------- 校验（信任边界：用户录入 / 上传文件） ----------

def _norm_period(raw) -> str:
    """归一化期间为 YYYY-MM；支持 datetime / '2026-09' / '2026/9'"""
    if isinstance(raw, datetime):
        return f"{raw.year:04d}-{raw.month:02d}"
    s = str(raw).strip().replace("/", "-")
    parts = s.split("-")
    if len(parts) < 2:
        raise ValueError(f"期间格式非法：{raw}")
    y, m = int(parts[0]), int(parts[1])
    if not (1 <= m <= 12):
        raise ValueError(f"月份非法：{raw}")
    return f"{y:04d}-{m:02d}"


def _valid_units(raw) -> Decimal:
    try:
        v = Decimal(str(raw))
    except (InvalidOperation, ValueError):
        raise ValueError(f"销量非数值：{raw}")
    if v < 0:
        raise ValueError(f"销量不可为负：{raw}")
    return v


def _upsert(db: Session, items: list[tuple], source: str) -> int:
    """items=[(period, channel, units:Decimal, note)] → upsert（键 period+channel）"""
    n = 0
    for period, channel, units, note in items:
        ex = (db.query(SalesActual)
              .filter(SalesActual.period == period, SalesActual.channel == channel).first())
        if ex:
            ex.units, ex.source = str(units), source
            if note is not None:
                ex.note = note
        else:
            db.add(SalesActual(period=period, channel=channel, units=str(units),
                               source=source, note=note))
        n += 1
    db.commit()
    return n


# ---------- 历史 ----------

@router.get("/actuals")
def list_actuals(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    """历史销量列表（按期间→渠道），并附总量与有效历史月数"""
    actuals = fs.load_actuals(db)
    total = fs._total_actual(actuals)
    rows = [{"period": a.period, "channel": a.channel, "units": a.units,
             "source": a.source, "note": a.note}
            for a in db.query(SalesActual).order_by(SalesActual.period,
                                                    SalesActual.channel).all()]
    return {"actuals": rows,
            "n_effective": fs.effective_months(total),
            "total_by_period": {p: str(v) for p, v in sorted(total.items())}}


@router.post("/actuals")
def upsert_actuals(body: ActualBatch, db: Session = Depends(get_db),
                   current_user=Depends(get_current_user)):
    """手工 upsert 历史销量（admin/editor）"""
    PermissionChecker.require_edit(current_user)
    try:
        items = [(_norm_period(a.period), a.channel, _valid_units(a.units), a.note)
                 for a in body.actuals]
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    bad = [c for _, c, _, _ in items if c not in _CHANNELS]
    if bad:
        raise HTTPException(status_code=400, detail=f"渠道非法：{bad}（须 online/offline）")
    n = _upsert(db, items, source="manual")
    OperationLogService.log(db, action="forecast.actuals.upsert",
                            description=f"录入历史销量 {n} 条", user_id=current_user.id)
    return {"ok": True, "upserted": n}


@router.get("/actuals/template")
def actuals_template(current_user=Depends(get_current_user)):
    """下载历史销量导入模板（.xlsx）——表头与 /actuals/import 解析口径一致"""
    import io
    import openpyxl
    from fastapi.responses import StreamingResponse

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "历史销量"
    ws.append(["月份", "线上(万台)", "线下(万台)"])
    # 两行示例：填真实数据时整表覆盖即可，重复期间按 period+channel 覆盖更新
    ws.append(["2026-01", 3.5, 1.2])
    ws.append(["2026-02", 3.8, 1.3])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="actuals_template.xlsx"'},
    )


@router.post("/actuals/import")
def import_actuals(file: UploadFile = File(...), db: Session = Depends(get_db),
                   current_user=Depends(get_current_user)):
    """Excel(.xlsx) 导入历史销量（admin/editor）——列：月份 | 线上(万台) | 线下(万台)"""
    PermissionChecker.require_edit(current_user)
    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 文件")
    import openpyxl
    try:
        wb = openpyxl.load_workbook(file.file, data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"无法解析 Excel：{e}")
    ws = wb.active
    items = []
    try:
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not row or row[0] in (None, "", "月份"):
                continue  # 跳过表头/空行
            period = _norm_period(row[0])
            if len(row) > 1 and row[1] not in (None, ""):
                items.append((period, "online", _valid_units(row[1]), "import"))
            if len(row) > 2 and row[2] not in (None, ""):
                items.append((period, "offline", _valid_units(row[2]), "import"))
    except ValueError as e:
        raise HTTPException(status_code=400, detail=f"第 {i + 1} 行：{e}")
    if not items:
        raise HTTPException(status_code=400, detail="未解析到有效数据行")
    n = _upsert(db, items, source="import")
    OperationLogService.log(db, action="forecast.actuals.import",
                            description=f"导入历史销量 {n} 条（{file.filename}）",
                            user_id=current_user.id)
    return {"ok": True, "upserted": n}


# ---------- 运行 / 接受 ----------

def _ser(res: dict) -> dict:
    """序列 Decimal → str（对齐 Cell.value 口径，避免 float 精度丢失）"""
    res = dict(res)
    res["series"] = [{k: (str(v) if isinstance(v, Decimal) else v)
                      for k, v in pt.items()} for pt in res["series"]]
    return res


@router.post("/run")
def run_forecast(body: RunRequest, db: Session = Depends(get_db),
                 current_user=Depends(get_current_user)):
    """运行预测（无副作用）：序列 + CI + 各档回测（全部角色可运行）"""
    horizon = tuple(body.horizon) if body.horizon else ("2026-08", "2027-12")
    try:
        res = fs.run_forecast(db, body.scenario_id, method=body.method,
                              horizon=horizon, scale_targets=body.scale_targets)
    except ValueError as e:
        if str(e) in ("no_version", "scenario_not_found"):
            raise HTTPException(status_code=404, detail="该情景尚未计算，请先导入或重算")
        raise HTTPException(status_code=400, detail=f"预测失败：{e}")
    return _ser(res)


@router.post("/apply")
def apply_forecast(body: ApplyRequest, db: Session = Depends(get_db),
                   current_user=Depends(get_current_user)):
    """接受预测 → 回填 qty + recalc（admin/editor）；给 lower/upper 则生成三情景"""
    PermissionChecker.require_edit(current_user)
    base = [p.model_dump() for p in body.base]
    lower = [p.model_dump() for p in body.lower] if body.lower else None
    upper = [p.model_dump() for p in body.upper] if body.upper else None
    if not base:
        raise HTTPException(status_code=400, detail="base 序列为空")
    try:
        out = fs.apply_forecast(db, body.scenario_id, base=base, method=body.method,
                                comment=body.comment, user_id=current_user.id,
                                lower=lower, upper=upper)
    except ValueError as e:
        if str(e) in ("no_version", "scenario_not_found"):
            raise HTTPException(status_code=404, detail="情景不存在或尚未计算")
        if str(e) in ("version_conflict_retry", "db_locked_retry"):
            raise HTTPException(status_code=409, detail="并发重算冲突，请重试")
        raise HTTPException(status_code=400, detail=f"应用预测失败：{e}")
    OperationLogService.log(db, action="forecast.apply",
                            description=f"情景 {body.scenario_id} 接受预测 → 版本 {out['base_version']}"
                                        + ("（三情景）" if lower and upper else ""),
                            detail=out, user_id=current_user.id)
    return out


@router.get("/runs")
def list_runs(scenario_id: int, db: Session = Depends(get_db),
              current_user=Depends(get_current_user)):
    """预测运行记录列表（最新在前）——回测报告 / 三情景对比用"""
    runs = (db.query(ForecastRun).filter(ForecastRun.scenario_id == scenario_id)
            .order_by(ForecastRun.id.desc()).all())
    return {"runs": [
        {"id": r.id, "method": r.method, "n_actual": r.n_actual,
         "horizon_start": r.horizon_start, "horizon_end": r.horizon_end,
         "base_version": r.base_version, "lower_version": r.lower_version,
         "upper_version": r.upper_version,
         "metrics": json.loads(r.metrics_json) if r.metrics_json else None,
         "created_at": r.created_at.isoformat() if r.created_at else None}
        for r in runs]}
